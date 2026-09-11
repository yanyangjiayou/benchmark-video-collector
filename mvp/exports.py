from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _collection_time(path: Path) -> datetime:
    parts = path.stem.split("_")
    if len(parts) >= 3:
        try:
            return datetime.strptime(f"{parts[1]}_{parts[2]}", "%Y%m%d_%H%M%S").astimezone()
        except ValueError:
            pass
    return datetime.fromtimestamp(path.stat().st_mtime).astimezone()


def _workbook_summary(path: Path) -> tuple[int, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        account_index = next((headers.index(name) for name in ("博主账号", "账号") if name in headers), None)
        accounts: list[str] = []
        count = 0
        for row in rows:
            if not any(value is not None for value in row):
                continue
            count += 1
            if account_index is not None and account_index < len(row):
                account = str(row[account_index] or "").strip()
                if account and account not in accounts:
                    accounts.append(account)
        return count, accounts
    finally:
        workbook.close()


def save_export_metadata(path: Path, *, accounts: list[str], account_label: str = "",
                         creator_url: str = "", keywords: list[str] | None = None) -> None:
    """Keep batch identity even when the workbook contains no new rows."""
    metadata = {"accounts": list(dict.fromkeys(account for account in accounts if account)),
                "account_label": account_label, "creator_url": creator_url,
                "keywords": keywords or []}
    destination = path.with_suffix(".json")
    temporary = destination.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(destination)


def _metadata(path: Path) -> dict:
    try:
        value = json.loads(path.with_suffix(".json").read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except (OSError, ValueError):
        return {}


def list_exports(output_dir: str | Path) -> list[dict[str, object]]:
    directory = Path(output_dir)
    if not directory.exists():
        return []
    exports = []
    for path in directory.glob("*.xlsx"):
        if path.name == "当前已完成结果.xlsx" or not path.is_file():
            continue
        platform = path.name.split("_", 1)[0]
        try:
            count, accounts = _workbook_summary(path)
            collected_at = _collection_time(path)
            size = path.stat().st_size
        except (OSError, ValueError, BadZipFile):
            continue
        metadata = _metadata(path)
        if not accounts:
            saved_accounts = metadata.get("accounts", [])
            if isinstance(saved_accounts, list):
                accounts = [str(account) for account in saved_accounts if account]
        if not accounts and metadata.get("account_label"):
            accounts = [str(metadata["account_label"])]
        exports.append({
            "filename": path.name,
            "platform": platform,
            "platform_name": {"xhs": "小红书", "dy": "抖音"}.get(platform, platform),
            "collected_at": collected_at.isoformat(timespec="seconds"),
            "collected_at_display": collected_at.strftime("%Y-%m-%d %H:%M:%S"),
            "row_count": count,
            "accounts": accounts,
            "account_display": "、".join(accounts) if accounts else "未记录账号",
            "creator_url": str(metadata.get("creator_url") or ""),
            "size_bytes": size,
        })
    return sorted(exports, key=lambda item: str(item["collected_at"]), reverse=True)


def resolve_export(output_dir: str | Path, filename: str) -> Path | None:
    if Path(filename).name != filename or not filename.lower().endswith(".xlsx"):
        return None
    directory = Path(output_dir).resolve()
    candidate = (directory / filename).resolve()
    if candidate.parent != directory or not candidate.is_file():
        return None
    return candidate
