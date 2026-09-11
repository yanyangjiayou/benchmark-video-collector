from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse


class CollectionHistory:
    """Persistent record of successfully collected platform content."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.path, timeout=30)
        connection.row_factory = sqlite3.Row
        return connection

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS collected_items (
                    platform TEXT NOT NULL,
                    content_id TEXT NOT NULL,
                    collection_key TEXT NOT NULL,
                    source_key TEXT NOT NULL DEFAULT '',
                    title TEXT NOT NULL DEFAULT '',
                    original_url TEXT NOT NULL DEFAULT '',
                    media_sha256 TEXT NOT NULL DEFAULT '',
                    batch_id TEXT NOT NULL DEFAULT '',
                    collected_at TEXT NOT NULL,
                    PRIMARY KEY (platform, content_id)
                )
                """
            )
            connection.execute(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS collected_items_media_sha256
                ON collected_items(media_sha256)
                WHERE media_sha256 <> ''
                """
            )

    @staticmethod
    def collection_key(platform: str, content_id: str) -> str:
        return f"{platform}:{content_id}"

    def seen_content_ids(self, platform: str, content_ids: Iterable[str]) -> set[str]:
        values = sorted({str(value) for value in content_ids if str(value)})
        if not values:
            return set()
        found: set[str] = set()
        with self._connect() as connection:
            for offset in range(0, len(values), 400):
                chunk = values[offset : offset + 400]
                placeholders = ",".join("?" for _ in chunk)
                rows = connection.execute(
                    f"SELECT content_id FROM collected_items WHERE platform = ? AND content_id IN ({placeholders})",
                    [platform, *chunk],
                )
                found.update(str(row["content_id"]) for row in rows)
        return found

    def seen_media_hashes(self, hashes: Iterable[str]) -> set[str]:
        values = sorted({str(value) for value in hashes if str(value)})
        if not values:
            return set()
        found: set[str] = set()
        with self._connect() as connection:
            for offset in range(0, len(values), 400):
                chunk = values[offset : offset + 400]
                placeholders = ",".join("?" for _ in chunk)
                rows = connection.execute(
                    f"SELECT media_sha256 FROM collected_items WHERE media_sha256 IN ({placeholders})",
                    chunk,
                )
                found.update(str(row["media_sha256"]) for row in rows)
        return found

    def record_many(self, entries: Iterable[dict[str, str]]) -> int:
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        before = 0
        after = 0
        with self._connect() as connection:
            before = connection.total_changes
            connection.executemany(
                """
                INSERT OR IGNORE INTO collected_items (
                    platform, content_id, collection_key, source_key, title,
                    original_url, media_sha256, batch_id, collected_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        entry["platform"],
                        entry["content_id"],
                        self.collection_key(entry["platform"], entry["content_id"]),
                        entry.get("source_key", ""),
                        entry.get("title", ""),
                        entry.get("original_url", ""),
                        entry.get("media_sha256", ""),
                        entry.get("batch_id", ""),
                        entry.get("collected_at", now),
                    )
                    for entry in entries
                ],
            )
            after = connection.total_changes
        return after - before

    def import_excel_exports(self, output_dir: str | Path) -> int:
        """Import successful rows from exports created before history tracking existed."""
        from openpyxl import load_workbook

        entries: list[dict[str, str]] = []
        for path in sorted(Path(output_dir).glob("*.xlsx")):
            try:
                workbook = load_workbook(path, read_only=True, data_only=True)
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = [str(value or "") for value in next(rows, ())]
                for values in rows:
                    record = dict(zip(headers, values))
                    transcript = str(record.get("视频转文字") or "")
                    if not transcript or transcript.startswith(("未取得", "转写失败")):
                        continue
                    marker = str(record.get("采集标记") or "")
                    platform, item_id = self._marker_parts(marker)
                    original_url = str(record.get("原始链接") or "")
                    if not platform or not item_id:
                        platform = path.name.split("_", 1)[0]
                        item_id = self._content_id_from_url(platform, original_url)
                    if platform in {"xhs", "dy"} and item_id:
                        entries.append({
                            "platform": platform,
                            "content_id": item_id,
                            "title": str(record.get("标题") or ""),
                            "original_url": original_url,
                            "batch_id": f"legacy:{path.stem}",
                        })
                workbook.close()
            except Exception:
                continue
        return self.record_many(entries)

    @staticmethod
    def _marker_parts(marker: str) -> tuple[str, str]:
        if ":" not in marker:
            return "", ""
        platform, item_id = marker.split(":", 1)
        return platform, item_id

    @staticmethod
    def _content_id_from_url(platform: str, url: str) -> str:
        parts = [part for part in urlparse(url).path.split("/") if part]
        if platform == "xhs" and "explore" in parts:
            index = parts.index("explore")
            return parts[index + 1] if index + 1 < len(parts) else ""
        if platform == "dy" and "video" in parts:
            index = parts.index("video")
            return parts[index + 1] if index + 1 < len(parts) else ""
        return ""
