from __future__ import annotations
import hashlib, json, os, re, shutil, subprocess, sys, time, uuid
from datetime import datetime
from pathlib import Path
from typing import Callable
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .exports import save_export_metadata
from .history import CollectionHistory
from .rules import CollectionRequest
from .sources import normalize_creator_url
from .transcriber import LocalTranscriber
from .xhs_video import download_missing_videos

ROOT = Path(__file__).resolve().parents[1]
VENDOR = ROOT / "vendor" / "MediaCrawler"


class JobCancelled(RuntimeError):
    """Raised when the user stops an active collection."""

def resolve_creator_url(url: str, platform: str) -> str:
    url = normalize_creator_url(url, platform)
    if not any(host in url for host in ("xhslink.cn", "xhslink.com", "v.douyin.com")):
        return url
    from urllib.request import Request, urlopen
    request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urlopen(request, timeout=20) as response:
            return normalize_creator_url(response.geturl(), platform, allow_short=False)
    except OSError as exc:
        raise ValueError("分享短链接解析失败，请在浏览器打开博主主页后，复制地址栏中的完整链接") from exc

def number(value: object) -> int:
    text = str(value or "0").strip().lower().replace(",", "")
    multiple = 10000 if ("万" in text or "w" in text) else 1
    try:
        return int(float(text.replace("万", "").replace("w", "")) * multiple)
    except ValueError:
        return 0

def published(item: dict, platform: str) -> datetime:
    stamp = float(item.get("time" if platform == "xhs" else "create_time") or 0)
    return datetime.fromtimestamp(stamp / 1000 if stamp > 10_000_000_000 else stamp)

def is_video(item: dict, platform: str) -> bool:
    return item.get("type") == "video" if platform == "xhs" else bool(item.get("video_download_url")) and not item.get("note_download_url")

def content_id(item: dict, platform: str) -> str:
    return str(item.get("note_id" if platform == "xhs" else "aweme_id") or "")

def media_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

def source_key(request: CollectionRequest) -> str:
    if request.trigger_type == "creator_url":
        return request.creator_url.strip()
    return ",".join(sorted(keyword.strip() for keyword in request.keywords if keyword.strip()))

def select_new_candidates(
    eligible_rows: list[dict], platform: str, max_items: int, history: CollectionHistory
) -> tuple[list[dict], dict[str, int]]:
    seen_ids = history.seen_content_ids(platform, (content_id(item, platform) for item in eligible_rows))
    candidates: list[dict] = []
    batch_ids: set[str] = set()
    history_skipped = 0
    duplicates_in_scan = 0
    missing_ids = 0
    for item in eligible_rows:
        item_id = content_id(item, platform)
        if not item_id:
            missing_ids += 1
            continue
        if item_id in batch_ids:
            duplicates_in_scan += 1
            continue
        batch_ids.add(item_id)
        if item_id in seen_ids:
            history_skipped += 1
            continue
        if len(candidates) < max_items:
            candidates.append(item)
    return candidates, {
        "eligible_total": len(batch_ids),
        "history_skipped": history_skipped,
        "duplicates_in_scan": duplicates_in_scan,
        "missing_ids": missing_ids,
    }

def find_media(raw: Path, platform: str, item_id: str) -> Path | None:
    for name in (["douyin", "dy"] if platform == "dy" else [platform]):
        folder = raw / name / "videos" / item_id
        if folder.exists():
            media = next((p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in {".mp4", ".webm", ".mov", ".m4a"}), None)
            if media:
                return media
    return None


def read_rows(raw: Path, platform: str) -> list[dict]:
    files = sorted(path for name in (["douyin", "dy"] if platform == "dy" else [platform])
                   for path in (raw / name / "jsonl").glob("*_contents_*.jsonl"))
    return [json.loads(line) for path in files for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]

def export_excel(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "视频结果"
    headers = ["采集标记", "博主账号", "标题", "发布时间", "视频转文字", "点赞数", "收藏数", "原始链接"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315EFB")
    for record in rows:
        ws.append([record[h] for h in headers])
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    for index, width in enumerate([30, 18, 38, 20, 80, 12, 12, 55], 1):
        ws.column_dimensions[chr(64 + index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[-1].hyperlink, row[-1].style = row[-1].value, "Hyperlink"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

def run_job(request: CollectionRequest, account_label: str, log: Callable[[str], None],
            progress: Callable[[dict], None] | None = None,
            on_records: Callable[[list[dict]], None] | None = None,
            model_name: str = "small", glossary: str = "",
            cancel_event: object | None = None,
            process_observer: Callable[[subprocess.Popen | None], None] | None = None,
            ) -> tuple[Path, dict, list[dict]]:
    job_id = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:6]
    raw = ROOT / "runtime" / "jobs" / job_id / "raw"
    raw.mkdir(parents=True)
    (raw.parent / "request.json").write_text(
        json.dumps({"request": request.model_dump(mode="json"), "account_label": account_label},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    history = CollectionHistory(ROOT / "runtime" / "collection_history.sqlite3")
    imported = history.import_excel_exports(ROOT / "output")
    if imported:
        log(f"已从旧 Excel 结果导入 {imported} 条历史采集标记")
    scan_limit = min(50, max(20, request.max_items * 4))
    cmd = [sys.executable, "main.py", "--platform", request.platform, "--lt", "qrcode",
           "--save_data_option", "jsonl", "--save_data_path", str(raw), "--crawler_max_notes_count", str(scan_limit),
           "--max_concurrency_num", "1", "--get_comment", "false", "--get_sub_comment", "false",
           "--enable_ip_proxy", "false"]
    if request.trigger_type == "creator_url":
        cmd += ["--type", "creator", "--creator_id", resolve_creator_url(request.creator_url, request.platform)]
    else:
        cmd += ["--type", "search", "--keywords", ",".join(request.keywords)]
    if request.platform == "xhs":
        cmd[1:2] = ["-m", "mvp.xhs_worker"]
    elif request.platform == "dy":
        cmd[1:2] = ["-m", "mvp.crawler_worker"]
    if request.platform == "xhs":
        log(
            f"正在按列表指标自动预筛选，最多扫描 {scan_limit} 条；"
            "只复用已确认登录的小红书窗口，不会在采集过程中打开、关闭或重新登录。"
        )
    else:
        log("正在使用专用浏览器配置检查登录会话；仅在登录失效时需要扫码。")
    env = {**os.environ, "PYTHONPATH": str(ROOT), "MPLCONFIGDIR": str(ROOT / "runtime/matplotlib"), "UV_CACHE_DIR": str(ROOT / "runtime/uv-cache")}
    process = subprocess.Popen(cmd, cwd=VENDOR, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                               text=True, encoding="utf-8", errors="replace")
    if process_observer:
        process_observer(process)
    errors: list[str] = []
    try:
        with (raw.parent / "crawler.log").open("w", encoding="utf-8") as logfile:
            for line in iter(process.stdout.readline, ""):
                if line.strip():
                    clean_line = re.sub(r"(?i)(cookie|authorization|xsec_token|mstoken)(['\"]?\s*[:=]\s*)[^\n]+", r"\1\2[已隐藏]", line.strip())
                    logfile.write(clean_line + "\n")
                    logfile.flush()
                    log(clean_line)
                    if any(marker in clean_line for marker in (
                        "ERROR", "Error:", "Exception:", "login qrcode not found",
                        "MVP_DOUYIN_LOGIN_REQUIRED", "MVP_XHS_",
                    )):
                        errors.append(clean_line)
        code = process.wait()
    finally:
        if process_observer:
            process_observer(None)
    if cancel_event is not None and getattr(cancel_event, "is_set")():
        raise JobCancelled("采集已由用户停止")
    worker_stats: dict = {}
    stats_path = raw.parent / "xhs_scan_summary.json"
    if request.platform == "xhs" and stats_path.is_file():
        try:
            worker_stats = json.loads(stats_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            worker_stats = {}
    if code:
        if worker_stats and progress:
            progress({
                "stage": "采集阶段已停止",
                "found": int(worker_stats.get("found", 0)),
                "details_requested": int(worker_stats.get("details_requested", 0)),
                "non_video_filtered": int(worker_stats.get("non_video_filtered", 0)),
                "likes_filtered": int(worker_stats.get("likes_filtered", 0)),
                "date_filtered": int(worker_stats.get("date_filtered", 0)),
                "date_prefiltered": int(worker_stats.get("date_prefiltered", 0)),
                "date_detail_filtered": int(worker_stats.get("date_detail_filtered", 0)),
                "followers_filtered": int(worker_stats.get("followers_filtered", 0)),
                "author_cards_filtered": int(worker_stats.get("author_cards_filtered", 0)),
                "creators_checked": int(worker_stats.get("creators_checked", 0)),
                "follower_filter_enabled": bool(worker_stats.get("follower_filter_enabled", False)),
                "trigger_type": worker_stats.get("trigger_type", request.trigger_type),
                "history_skipped": int(worker_stats.get("history_skipped", 0)),
                "metric_missing_skipped": sum(
                    int(worker_stats.get(key, 0))
                    for key in ("likes_missing", "time_missing", "followers_missing")
                ),
                "eligible": int(worker_stats.get("selected", 0)),
            })
        if any("MVP_DOUYIN_LOGIN_REQUIRED" in line for line in errors):
            raise RuntimeError("抖音登录已失效，请点击“确认登录”并在抖音窗口重新扫码")
        if any("MVP_XHS_ACCESS_RESTRICTED" in line for line in errors):
            raise RuntimeError("小红书返回操作频繁、安全验证或访问限制；任务已立即停止，且没有自动重试或重新登录")
        if any("MVP_XHS_LOGIN_REQUIRED" in line for line in errors):
            raise RuntimeError("小红书登录态未就绪；任务已停止，未在后台自动打开登录页。请先点击「确认登录」并完成扫码，保持浏览器窗口不关闭，再重新采集")
        if any("MVP_XHS_METRIC_UNAVAILABLE" in line for line in errors):
            raise RuntimeError("小红书未返回可核验的点赞数、粉丝数或发布时间；为避免错误筛选，本次没有把不确定内容导出")
        if any("MVP_XHS_LOGIN_CHECK_FAILED" in line for line in errors):
            raise RuntimeError("小红书登录状态检查请求异常；任务已停止且没有自动重新登录，请稍后再试")
        raise RuntimeError(f"采集失败（退出代码 {code}），请查看运行记录。详细日志已保存在本机。")
    all_rows = read_rows(raw, request.platform)
    if not all_rows and errors:
        raise RuntimeError("平台未返回内容，运行记录中有登录或请求错误；请确认主页链接，并重新确认登录后重试。")
    video_rows = [x for x in all_rows if is_video(x, request.platform)]
    date_rows = [x for x in video_rows if request.start_date <= published(x, request.platform).date() <= request.end_date]
    eligible_rows = sorted(
        (x for x in date_rows if number(x.get("liked_count")) >= request.min_likes),
        key=lambda x: published(x, request.platform),
        reverse=True,
    )
    candidates, selection = select_new_candidates(eligible_rows, request.platform, request.max_items, history)
    history_skipped = selection["history_skipped"] + int(worker_stats.get("history_skipped", 0))
    duplicates_in_scan = selection["duplicates_in_scan"] + int(worker_stats.get("duplicates_in_scan", 0))
    likes_missing = int(worker_stats.get("likes_missing", 0))
    time_missing = int(worker_stats.get("time_missing", 0))
    followers_missing = int(worker_stats.get("followers_missing", 0))
    metric_missing_skipped = likes_missing + time_missing + followers_missing
    summary = {"found": int(worker_stats.get("found", len(all_rows))),
               "videos": int(worker_stats.get("videos", len(video_rows))),
               "non_video_filtered": int(worker_stats.get("non_video_filtered", len(all_rows) - len(video_rows))),
               "in_date_range": int(worker_stats.get("in_date_range", len(date_rows))),
               "date_filtered": int(worker_stats.get("date_filtered", len(video_rows) - len(date_rows))),
               "eligible_total": selection["eligible_total"], "eligible": len(candidates), "history_skipped": history_skipped,
               "duplicates_in_scan": duplicates_in_scan, "content_duplicates": 0,
               "likes_filtered": int(worker_stats.get("likes_filtered", 0)),
               "followers_filtered": int(worker_stats.get("followers_filtered", 0)),
               "author_cards_filtered": int(worker_stats.get("author_cards_filtered", 0)),
               "creators_checked": int(worker_stats.get("creators_checked", 0)),
               "follower_filter_enabled": bool(worker_stats.get("follower_filter_enabled", request.min_followers is not None)),
               "trigger_type": worker_stats.get("trigger_type", request.trigger_type),
               "date_prefiltered": int(worker_stats.get("date_prefiltered", 0)),
               "date_detail_filtered": int(worker_stats.get("date_detail_filtered", 0)),
               "likes_missing": likes_missing,
               "time_missing": time_missing,
               "followers_missing": followers_missing,
               "metric_missing_skipped": metric_missing_skipped,
               "details_requested": int(worker_stats.get("details_requested", len(all_rows) if request.platform == "xhs" else 0)),
               "profile_requests": int(worker_stats.get("profile_requests", 0)),
               "list_requests": int(worker_stats.get("list_requests", 0)),
               "transcribed": 0, "media_missing": 0, "transcribe_errors": 0, "exported": 0,
               "stage": "筛选完成", "estimated_total_seconds": 90 + len(candidates) * (75 if model_name == "small" else 130)}
    if progress:
        progress(summary)
    log(
        f"筛选结果：读取 {summary['found']} 条列表卡片，排除 {summary['non_video_filtered']} 条非视频，"
        f"对 {summary['details_requested']} 条候选读取详情核验，最终新采集 {len(candidates)} 条"
    )
    if summary["date_prefiltered"]:
        log(f"列表日期预筛已排除 {summary['date_prefiltered']} 条，无需读取这些详情")
    if summary["date_detail_filtered"]:
        log(f"详情日期复核后排除 {summary['date_detail_filtered']} 条内容")
    if summary["likes_filtered"]:
        log(f"列表点赞数预筛选已排除 {summary['likes_filtered']} 条，无需打开这些详情")
    if summary["followers_filtered"]:
        log(f"粉丝数筛选已排除 {summary['followers_filtered']} 位博主")
    if metric_missing_skipped:
        log(
            f"有 {metric_missing_skipped} 条因平台未返回可核验指标而跳过"
            f"（点赞 {likes_missing}、发布时间 {time_missing}、粉丝 {followers_missing}）；"
            "这些不确定内容没有导出，任务继续处理其余内容"
        )
    if history_skipped:
        log(f"已根据持久化采集标记跳过 {history_skipped} 条历史内容")
    if duplicates_in_scan:
        log(f"扫描结果内去除 {duplicates_in_scan} 条重复内容")
    if selection["missing_ids"]:
        log(f"已跳过 {selection['missing_ids']} 条缺少内容 ID 的记录")
    if request.platform == "xhs" and candidates:
        if cancel_event is not None and getattr(cancel_event, "is_set")():
            raise JobCancelled("采集已由用户停止")
        missing = [item for item in candidates if find_media(raw, request.platform, content_id(item, request.platform)) is None]
        if missing:
            summary["stage"] = "正在取得视频文件"
            if progress:
                progress(summary)
            log(f"MediaCrawler 未取得 {len(missing)} 个视频文件，正在从详情页补取。")
            download_missing_videos(missing, raw, log)
    prepared: list[tuple[dict, Path | None, str]] = []
    hashes = {}
    for item in candidates:
        item_id = content_id(item, request.platform)
        media = find_media(raw, request.platform, item_id)
        file_hash = media_sha256(media) if media else ""
        hashes[item_id] = file_hash
    historical_hashes = history.seen_media_hashes(hashes.values())
    batch_hashes: set[str] = set()
    for item in candidates:
        item_id = content_id(item, request.platform)
        media = find_media(raw, request.platform, item_id)
        file_hash = hashes[item_id]
        if file_hash and (file_hash in historical_hashes or file_hash in batch_hashes):
            summary["content_duplicates"] += 1
            log(f"视频内容与已采集项完全相同，跳过：{item.get('title', '')[:30]}")
            continue
        if file_hash:
            batch_hashes.add(file_hash)
        prepared.append((item, media, file_hash))
    summary["eligible"] = len(prepared)
    if progress:
        progress(summary)

    def model_stage(message: str) -> None:
        summary["stage"] = message
        log(message)
        if progress:
            progress(summary)

    transcriber = LocalTranscriber(model_name, on_stage=model_stage)
    if request.transcribe_video and any(media for _, media, _ in prepared):
        model_started = time.monotonic()
        transcriber.load()
        summary["model_load_seconds"] = round(time.monotonic() - model_started, 1)
    results: list[dict] = []
    history_entries: list[dict[str, str]] = []
    for index, (item, media, file_hash) in enumerate(prepared, 1):
        if cancel_event is not None and getattr(cancel_event, "is_set")():
            raise JobCancelled("采集已由用户停止")
        item_id = content_id(item, request.platform)
        log(f"正在转写 {index}/{len(prepared)}：{item.get('title', '')[:30]}")
        summary.update(stage=f"正在转写第 {index}/{len(prepared)} 条", current=index)
        if progress:
            progress(summary)
        completed = False
        if media:
            if request.transcribe_video:
                prompt = "，".join(x for x in [glossary.strip(), account_label.strip(), item.get("title", ""), item.get("desc", "")] if x)[:1000]
                try:
                    transcript = transcriber.transcribe(media, initial_prompt=prompt)
                    summary["transcribed"] += 1
                    completed = True
                except Exception as exc:
                    transcript = f"转写失败：{exc}"
                    summary["transcribe_errors"] += 1
                    log(f"转写失败：{item.get('title', '')[:30]}（{exc}）")
            else:
                transcript = "未启用视频转写"
                completed = True
        else:
            transcript = "未取得视频文件，无法转写"
            summary["media_missing"] += 1
        original_url = item.get("note_url") or item.get("aweme_url") or ""
        results.append({"采集标记": history.collection_key(request.platform, item_id),
                        "博主账号": account_label or item.get("nickname", ""), "标题": item.get("title", ""),
                        "发布时间": published(item, request.platform).strftime("%Y-%m-%d %H:%M:%S"),
                        "视频转文字": transcript, "点赞数": number(item.get("liked_count")),
                        "收藏数": number(item.get("collected_count")),
                        "原始链接": original_url})
        if completed:
            history_entries.append({"platform": request.platform, "content_id": item_id,
                                    "source_key": source_key(request), "title": str(item.get("title", "")),
                                    "original_url": str(original_url), "media_sha256": file_hash,
                                    "batch_id": job_id})
        if on_records:
            on_records(results)
        if progress:
            progress(summary)
    destination = ROOT / "output" / f"{request.platform}_{job_id}.xlsx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    accounts = list(dict.fromkeys(str(row["博主账号"]) for row in results if row.get("博主账号")))
    if not accounts and account_label:
        accounts = [account_label]
    if not accounts and request.trigger_type == "creator_url":
        accounts = list(dict.fromkeys(str(item["nickname"]) for item in all_rows if item.get("nickname")))
    save_export_metadata(destination, accounts=accounts, account_label=account_label,
                         creator_url=request.creator_url if request.trigger_type == "creator_url" else "",
                         keywords=request.keywords if request.trigger_type == "keyword" else [])
    export_excel(results, destination)
    summary["history_marked"] = history.record_many(history_entries)
    summary["exported"] = len(results)
    summary["stage"] = "全部完成"
    if progress:
        progress(summary)
    shutil.rmtree(raw, ignore_errors=True)
    log(f"已完成：{destination}（新增采集标记 {summary['history_marked']} 条）")
    return destination, summary, results
